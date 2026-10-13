"""Document ingestion for DOCX, PPTX, TXT, Markdown, CSV, and Excel.

Each format has a dedicated extractor that returns structure-aware markdown:
headings become `##` sections, tables become markdown pipe tables, and slide
or sheet boundaries become section markers. That structure flows into the AI
summary, semantic chunking (section metadata) and retrieval, the same way the
Docling PDF pipeline does for PDFs.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile

from app.core.config import get_settings
from app.core.logging import get_logger

logger = get_logger(__name__)

SUPPORTED = {".docx", ".pptx", ".txt", ".md", ".markdown", ".csv", ".xlsx"}

SUPPORTED_LABELS = "DOCX, PPTX, XLSX, CSV, TXT, Markdown (MD)"

_ZIP_EXTENSIONS = {".docx", ".pptx", ".xlsx"}


def _validate_zip_bomb(data: bytes, filename: str) -> None:
    """Reject ZIP-based files that exceed uncompressed size or ratio limits."""
    lower = filename.lower()
    ext = "." + lower.rsplit(".", 1)[-1] if "." in lower else ""
    if ext not in _ZIP_EXTENSIONS:
        return
    settings = get_settings()
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            total = 0
            for info in zf.infolist():
                total += info.file_size
                if total > settings.zip_max_uncompressed_bytes:
                    raise ValueError(
                        f"ZIP uncompressed size ({total} bytes) exceeds limit "
                        f"({settings.zip_max_uncompressed_bytes} bytes)"
                    )
            if settings.zip_max_uncompressed_bytes > 0 and len(data) > 0:
                ratio = total / len(data)
                if ratio > settings.zip_max_ratio:
                    raise ValueError(
                        f"ZIP decompression ratio ({ratio:.1f}:1) exceeds limit "
                        f"({settings.zip_max_ratio}:1)"
                    )
    except zipfile.BadZipFile:
        # Not a valid ZIP; let the format-specific extractor handle the error
        pass


def _md_escape(cell: str) -> str:
    return (cell or "").replace("|", "\\|").replace("\n", " ").strip()


def _md_table(rows: list[list[str]]) -> str:
    """Render rows as a markdown pipe table (first row = header)."""
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [list(_md_escape(c) for c in r) for r in rows]
    norm = [r + [""] * (width - len(r)) for r in norm]
    header = norm[0]
    body = norm[1:]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


def _is_heading(paragraph) -> bool:
    style = ""
    try:
        style = (paragraph.style.name or "") if paragraph.style else ""
    except Exception:
        return False
    return bool(re.match(r"^(Title|Heading\s*[1-9]|subtitle|TOC\s*Heading)", style, re.I))


def _heading_level(paragraph) -> int:
    style = (paragraph.style.name or "") if paragraph.style else ""
    match = re.search(r"([1-9])", style)
    if match and re.match(r"^Heading", style, re.I):
        return min(int(match.group(1)), 6)
    return 1


def _extract_docx(data: bytes) -> str:
    """DOCX -> markdown: headings stay headings, tables become pipe tables."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(io.BytesIO(data))
    parts: list[str] = []
    body = doc.element.body

    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, doc)
            text = paragraph.text.strip()
            if not text:
                continue
            if _is_heading(paragraph):
                level = _heading_level(paragraph)
                parts.append(f"{'#' * level} {text}")
            else:
                parts.append(text)
        elif child.tag.endswith("}tbl"):
            table = Table(child, doc)
            rows: list[list[str]] = []
            for row in table.rows:
                rows.append([cell.text for cell in row.cells])
            md = _md_table(rows)
            if md:
                parts.append(md)

    return "\n\n".join(parts).strip()


def _extract_pptx(data: bytes) -> str:
    """PPTX -> markdown: slide titles as headings, tables as pipe tables."""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_lines: list[str] = []
        title = None
        table_blocks: list[str] = []

        # Preserve top-to-bottom reading order
        shapes = sorted(
            slide.shapes,
            key=lambda s: (s.top if isinstance(s.top, Emu) else 0),
        )
        for shape in shapes:
            if shape.has_table:
                rows = [[c.text for c in row.cells] for row in shape.table.rows]
                md = _md_table(rows)
                if md:
                    table_blocks.append(md)
                continue
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            if title is None and shape == slide.shapes.title:
                title = text.split("\n")[0].strip()
            else:
                slide_lines.append(text)

        heading = title or f"Slide {i}"
        slide_lines_md = "\n\n".join(s for s in slide_lines if s)
        block = f"## {heading} (Slide {i})"
        if slide_lines_md:
            block += f"\n\n{slide_lines_md}"
        for table_md in table_blocks:
            block += f"\n\n{table_md}"
        parts.append(block)

    return "\n\n".join(parts).strip()


def _extract_xlsx(data: bytes) -> str:
    """XLSX -> markdown: one `## Sheet` section per worksheet, header table."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(c.strip() for c in cells):
                rows.append(cells)
        if not rows:
            continue
        block = f"## Sheet: {ws.title}"
        table = _md_table(rows)
        if table:
            block += f"\n\n{table}"
        parts.append(block)
    wb.close()
    return "\n\n".join(parts).strip()


def _extract_csv(data: bytes) -> str:
    """CSV -> markdown pipe table with the first row as header."""
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [list(row) for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        return ""
    return _md_table(rows)


def _extract_plain(data: bytes) -> str:
    return data.decode("utf-8", errors="replace").strip()


def extract_text(data: bytes, filename: str) -> dict:
    """Extract structure-aware text from a supported document.

    Returns {title, text} where `text` is markdown preserving headings,
    tables and section boundaries for the AI pipeline.
    """
    lower = filename.lower()
    ext = "." + lower.rsplit(".", 1)[-1] if "." in lower else ""
    if ext not in SUPPORTED:
        raise ValueError(
            f"Unsupported document type: {ext or filename}. "
            f"Supported formats: {SUPPORTED_LABELS}"
        )

    if ext in _ZIP_EXTENSIONS:
        _validate_zip_bomb(data, filename)

    if ext == ".docx":
        text = _extract_docx(data)
    elif ext == ".pptx":
        text = _extract_pptx(data)
    elif ext == ".xlsx":
        text = _extract_xlsx(data)
    elif ext == ".csv":
        text = _extract_csv(data)
    else:  # .txt, .md, .markdown
        text = _extract_plain(data)

    if not text:
        raise ValueError(f"No extractable text found in {filename}")

    return {"title": filename, "text": text}
